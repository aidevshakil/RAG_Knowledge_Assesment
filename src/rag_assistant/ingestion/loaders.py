"""Document loaders.

Each loader returns `LoadedDocument`: normalised text plus optional page
offsets, which is what lets citations point at a real PDF page. Loaders are
registered by extension, so supporting a new format is one function.
"""

from __future__ import annotations

import csv
import io
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path

from rag_assistant.core.exceptions import (
    DependencyMissingError,
    DocumentLoadError,
    UnsupportedFileTypeError,
)
from rag_assistant.core.logging_config import get_logger
from rag_assistant.core.text import normalize_whitespace
from rag_assistant.core.types import SourceType

logger = get_logger(__name__)


@dataclass(slots=True)
class LoadedDocument:
    """Extracted text plus enough structure to build citations."""

    text: str
    source_type: str
    page_offsets: list[tuple[int, int]] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)

    def page_for_offset(self, offset: int) -> int | None:
        """Which source page a character offset falls on (binary search)."""
        if not self.page_offsets:
            return None
        low, high = 0, len(self.page_offsets) - 1
        page = self.page_offsets[0][1]
        while low <= high:
            mid = (low + high) // 2
            start, number = self.page_offsets[mid]
            if start <= offset:
                page = number
                low = mid + 1
            else:
                high = mid - 1
        return page


Loader = Callable[[bytes, str], LoadedDocument]
LOADERS: dict[str, Loader] = {}


def register_loader(*extensions: str) -> Callable[[Loader], Loader]:
    def decorator(func: Loader) -> Loader:
        for ext in extensions:
            LOADERS[ext.lower().lstrip(".")] = func
        return func

    return decorator


def supported_extensions() -> list[str]:
    return sorted(LOADERS)


@register_loader("txt", "md", "markdown", "rst", "log", "json", "py")
def load_text(data: bytes, name: str) -> LoadedDocument:
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 never fails
        raise DocumentLoadError(f"Could not decode {name} as text")
    ext = Path(name).suffix.lstrip(".").lower()
    source_type = SourceType.MARKDOWN if ext in {"md", "markdown"} else SourceType.TXT
    return LoadedDocument(text=normalize_whitespace(text), source_type=str(source_type))


@register_loader("pdf")
def load_pdf(data: bytes, name: str) -> LoadedDocument:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DependencyMissingError("pypdf", "PDF ingestion") from exc

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise DocumentLoadError(f"{name} is not a readable PDF: {exc}") from exc

    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception as exc:
            raise DocumentLoadError(f"{name} is password protected") from exc

    parts: list[str] = []
    offsets: list[tuple[int, int]] = []
    cursor = 0
    for number, page in enumerate(reader.pages, start=1):
        try:
            raw = page.extract_text() or ""
        except Exception as exc:
            logger.warning("%s page %d: %s", name, number, exc)
            continue
        text = normalize_whitespace(raw)
        if not text:
            continue
        offsets.append((cursor, number))
        parts.append(text)
        cursor += len(text) + 2

    if not parts:
        raise DocumentLoadError(
            f"No extractable text in {name}. It is probably a scanned image — "
            "run OCR (e.g. ocrmypdf) before uploading."
        )
    return LoadedDocument(
        text="\n\n".join(parts),
        source_type=str(SourceType.PDF),
        page_offsets=offsets,
        metadata={"pages": len(reader.pages)},
    )


@register_loader("docx")
def load_docx(data: bytes, name: str) -> LoadedDocument:
    try:
        import docx
    except ImportError as exc:
        raise DependencyMissingError("python-docx", "DOCX ingestion") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentLoadError(f"{name} is not a readable .docx: {exc}") from exc

    blocks = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    text = normalize_whitespace("\n\n".join(blocks))
    if not text:
        raise DocumentLoadError(f"No text found in {name}")
    return LoadedDocument(text=text, source_type=str(SourceType.DOCX))


@register_loader("csv", "tsv")
def load_csv(data: bytes, name: str) -> LoadedDocument:
    text = data.decode("utf-8", errors="replace")
    delimiter = "\t" if name.lower().endswith(".tsv") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = list(reader)
    if not rows:
        raise DocumentLoadError(f"{name} is empty")
    header, *body = rows
    lines = [
        "; ".join(f"{key}: {value}" for key, value in zip(header, row, strict=False) if value)
        for row in body
    ]
    return LoadedDocument(
        text=normalize_whitespace("\n".join(line for line in lines if line)),
        source_type=str(SourceType.CSV),
        metadata={"rows": len(body), "columns": len(header)},
    )


@register_loader("html", "htm")
def load_html(data: bytes, name: str) -> LoadedDocument:
    text = _html_to_text(data.decode("utf-8", errors="replace"))
    if not text:
        raise DocumentLoadError(f"No text found in {name}")
    return LoadedDocument(text=text, source_type=str(SourceType.HTML))


_SCRIPT_STYLE = re.compile(r"<(script|style|noscript)[^>]*>.*?</\1>", re.S | re.I)
_TAGS = re.compile(r"<[^>]+>")


def _html_to_text(html: str) -> str:
    """BeautifulSoup when available, regex fallback otherwise."""
    try:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(html, "html.parser")
        for tag in soup(["script", "style", "noscript", "nav", "footer", "header"]):
            tag.decompose()
        return normalize_whitespace(soup.get_text("\n"))
    except ImportError:
        stripped = _TAGS.sub(" ", _SCRIPT_STYLE.sub(" ", html))
        import html as html_module

        return normalize_whitespace(html_module.unescape(stripped))


def load_bytes(data: bytes, name: str) -> LoadedDocument:
    ext = Path(name).suffix.lstrip(".").lower()
    loader = LOADERS.get(ext)
    if loader is None:
        raise UnsupportedFileTypeError(
            f"'.{ext or name}' is not supported. Supported: {', '.join(supported_extensions())}"
        )
    loaded = loader(data, name)
    if not loaded.text.strip():
        raise DocumentLoadError(f"{name} produced no text")
    return loaded


def load_path(path: str | Path) -> LoadedDocument:
    path = Path(path).expanduser()
    if not path.is_file():
        raise DocumentLoadError(f"{path} is not a file")
    return load_bytes(path.read_bytes(), path.name)


def load_url(url: str, *, timeout: float = 20.0) -> LoadedDocument:
    """Fetch a web page (or a linked PDF) and extract its text."""
    try:
        import requests
    except ImportError as exc:
        raise DependencyMissingError("requests", "URL ingestion") from exc
    try:
        response = requests.get(
            url, timeout=timeout, headers={"User-Agent": "rag-knowledge-assistant/1.0"}
        )
        response.raise_for_status()
    except Exception as exc:
        raise DocumentLoadError(f"Could not fetch {url}: {exc}") from exc

    content_type = response.headers.get("content-type", "").lower()
    if "pdf" in content_type or url.lower().endswith(".pdf"):
        loaded = load_pdf(response.content, url)
    else:
        loaded = LoadedDocument(text=_html_to_text(response.text), source_type=str(SourceType.URL))
    if not loaded.text.strip():
        raise DocumentLoadError(f"No readable text at {url}")
    loaded.metadata["url"] = url
    loaded.source_type = str(SourceType.URL)
    return loaded
